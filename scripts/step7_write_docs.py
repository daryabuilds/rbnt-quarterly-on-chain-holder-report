"""Step 7. Write the spreadsheet, the CSV bundle and the community summary.

Reads public/data/report.json and the CSV files from step 6, so every number
here is the same number shown on the site and in the PDF.

Outputs:
  public/data/RBNT-Holder-Report-<quarter>-data.xlsx
  public/data/RBNT-Holder-Report-<quarter>-csv.zip
  COMMUNITY_SUMMARY.md (and a copy in public/downloads)
  public/downloads/METHOD.md (copy)
  public/data/summary_figures.json (figures quoted in the summary, for verify_outputs.py)
"""

import csv
import glob
import json
import os
import shutil
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import PUBLIC_DATA, ROOT

SHEETS = {
    "01_where_rbnt_sits.csv": "Where RBNT sits",
    "02_holders_by_balance_band.csv": "Holders by band",
    "03_concentration.csv": "Concentration",
    "04_largest_holder_groups.csv": "Largest holders",
    "05_staking_pools.csv": "Staking pools",
    "06_staking_flows_in_quarter.csv": "Staking flows",
    "07_named_wallets.csv": "Named wallets",
    "08_named_wallet_flows_in_quarter.csv": "Named wallet flows",
    "09_wrbnt.csv": "WRBNT",
    "10_verification_checks.csv": "Checks",
}

QUOTED = []  # every figure written into the summary, checked by verify_outputs.py


def q(s):
    QUOTED.append(s)
    return s


def full(v, d=0):
    return q(f"{v:,.{d}f}")


def mil(v):
    return q(f"{v / 1e6:,.1f}") + " million"


def bil(v):
    return q(f"{v / 1e9:,.2f}") + " billion"


def amt(v):
    return bil(v) if abs(v) >= 1e9 else mil(v) if abs(v) >= 1e6 else full(v)


def pct(v, d=1):
    return q(f"{v * 100:.{d}f}") + "%"


def rel(a, c):
    return q(f"{abs((c - a) / a) * 100:.1f}") + "%"


def num(s):
    s = s.strip()
    if s == "":
        return None
    try:
        f = float(s.replace(",", ""))
        return int(f) if f.is_integer() and "." not in s else f
    except ValueError:
        return s


def workbook(rep, label):
    wb = Workbook()
    ws = wb.active
    ws.title = "About"
    head_fill = PatternFill("solid", fgColor="1D3140")
    head_font = Font(bold=True, color="F1F5F8")
    snaps = rep["meta"]["snapshots"]
    rows = [
        ["RBNT Quarterly On-Chain Holder Report, " + rep["meta"]["quarter"]["label"]],
        ["Period", "1 July 2026 00:00:00 UTC to 30 September 2026 23:59:59 UTC"],
        ["Chain", "Redbelly mainnet, chain ID 151"],
        ["Data verified", rep["meta"]["generated_utc"]],
        [],
        ["Snapshot", "State at (UTC)", "Block", "Block time (UTC)"],
    ] + [[s["key"], s["state_at_utc"], s["block"], s["block_time_utc"]] for s in snaps] + [
        [],
        ["Sources"],
        ["Balances, contract reads, blocks", "Redbelly mainnet RPC, https://governors.mainnet.redbelly.network"],
        ["Internal transfers, event logs, block by time", "Routescan API, https://api.routescan.io/v2/network/mainnet/evm/151/etherscan/api"],
        ["Reserved wallet names and allocations", "https://redbellydao.network/treasury/"],
        ["Explorer", "https://redbelly.routescan.io"],
        [],
        ["Notes"],
        ["Amounts are in RBNT, rounded to whole RBNT. Shares are percentages."],
        ["Holdings count RBNT in the wallet, WRBNT held, and RBNT staked in the Reddex pools, so each RBNT is counted once."],
        ["Only publicly named wallets appear by address. All other addresses are grouped."],
        ["Full method: METHOD.md"],
    ]
    for r in rows:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=14)
    for c in ws[7]:
        c.fill, c.font = head_fill, head_font
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 26

    for fname, title in SHEETS.items():
        path = os.path.join(PUBLIC_DATA, fname)
        sh = wb.create_sheet(title)
        with open(path, newline="") as f:
            data = list(csv.reader(f))
        for i, r in enumerate(data):
            sh.append(r if i == 0 else [num(x) for x in r])
        for c in sh[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(wrap_text=True, vertical="top")
        sh.row_dimensions[1].height = 48
        for col in range(1, len(data[0]) + 1):
            letter = get_column_letter(col)
            width = max(len(str(r[col - 1])) if col - 1 < len(r) else 0 for r in data[1:]) if len(data) > 1 else 10
            sh.column_dimensions[letter].width = min(60, max(14, width + 2))
            for cell in sh[letter][1:]:
                if isinstance(cell.value, int):
                    cell.number_format = "#,##0"
                elif isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"
        sh.freeze_panes = "B2"
    out = os.path.join(PUBLIC_DATA, f"RBNT-Holder-Report-{label}-data.xlsx")
    wb.save(out)
    return out


def summary(rep):
    from summary_text import build  # kept in its own file so the words are easy to edit
    return build(rep, {"full": full, "mil": mil, "bil": bil, "amt": amt, "pct": pct, "rel": rel, "q": q})


def main():
    rep = json.load(open(os.path.join(PUBLIC_DATA, "report.json")))
    label = rep["meta"]["quarter"]["label"].replace(" ", "-")
    print("wrote", workbook(rep, label))

    zpath = os.path.join(PUBLIC_DATA, f"RBNT-Holder-Report-{label}-csv.zip")
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(glob.glob(os.path.join(PUBLIC_DATA, "*.csv"))):
            z.write(f, os.path.basename(f))
        z.write(os.path.join(ROOT, "METHOD.md"), "METHOD.md")
    print("wrote", zpath)

    text = summary(rep)
    with open(os.path.join(ROOT, "COMMUNITY_SUMMARY.md"), "w") as f:
        f.write(text)
    os.makedirs(os.path.join(ROOT, "public", "downloads"), exist_ok=True)
    shutil.copy(os.path.join(ROOT, "COMMUNITY_SUMMARY.md"), os.path.join(ROOT, "public", "downloads", "COMMUNITY_SUMMARY.md"))
    shutil.copy(os.path.join(ROOT, "METHOD.md"), os.path.join(ROOT, "public", "downloads", "METHOD.md"))
    with open(os.path.join(PUBLIC_DATA, "summary_figures.json"), "w") as f:
        json.dump(sorted(set(QUOTED)), f, indent=0)
    print("wrote COMMUNITY_SUMMARY.md")


if __name__ == "__main__":
    main()
